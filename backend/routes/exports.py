"""Выгрузка отчётов ситуационного центра в PDF и Excel.

Вынесено из analytics.py: генерация документов — отдельная забота, и тянуть
reportlab с openpyxl в модуль с бизнес-логикой аналитики незачем. Префикс тот
же, поэтому для фронтенда это всё те же /api/analytics/export/*.
"""
import io
from datetime import datetime
from pathlib import Path

from fastapi import APIRouter
from fastapi.responses import StreamingResponse

from routes.analytics import regions, spikes, summary
from store import all_tickets

router = APIRouter(prefix="/api/analytics/export", tags=["export"])

# Показатели «до внедрения» — демонстрационные ориентиры, те же, что в UI.
BASELINE = {
    "avg_processing_time": "~5 мин",
    "resolved_percent": "~78%",
    "routing_errors": "~15%",
}


def _register_cyrillic_font() -> str:
    """Зарегистрировать TTF с поддержкой кириллицы.

    Встроенные шрифты reportlab (Helvetica и прочие Type1) кириллицу не
    отображают — вместо букв выходят чёрные квадраты. Берём системный Arial,
    на Linux — DejaVu. Если ничего не нашлось, возвращаем Helvetica: текст
    будет испорчен, но генерация не упадёт.
    """
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        r"C:\Windows\Fonts\arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/Library/Fonts/Arial.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            try:
                pdfmetrics.registerFont(TTFont("AURASans", path))
                return "AURASans"
            except Exception:
                continue
    return "Helvetica"


@router.get("/pdf")
def export_pdf():
    """PDF-отчёт: ключевые метрики, таблица регионов, активные всплески."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    font = _register_cyrillic_font()
    summary_data = summary()
    region_rows = regions()
    spike_rows = spikes()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="AURA Pulse 109 - отчёт",
    )

    base = getSampleStyleSheet()
    h1 = ParagraphStyle(
        "AuraH1", parent=base["Heading1"], fontName=font, fontSize=16,
        textColor=colors.HexColor("#111827"), spaceAfter=2,
    )
    h2 = ParagraphStyle(
        "AuraH2", parent=base["Heading2"], fontName=font, fontSize=11,
        textColor=colors.HexColor("#6B7280"), spaceBefore=16, spaceAfter=6,
    )
    body = ParagraphStyle(
        "AuraBody", parent=base["Normal"], fontName=font, fontSize=9,
        textColor=colors.HexColor("#111827"),
    )
    muted = ParagraphStyle(
        "AuraMuted", parent=body, fontSize=8,
        textColor=colors.HexColor("#9CA3AF"),
    )

    table_style = TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F9FAFB")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#6B7280")),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#E5E7EB")),
        ("GRID", (0, 1), (-1, -1), 0.3, colors.HexColor("#F3F4F6")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ])

    total_pretty = f"{summary_data.total_tickets:,}".replace(",", " ")
    story = [
        Paragraph("AURA - Pulse 109", h1),
        Paragraph(
            "Отчёт ситуационного центра · сформирован "
            + datetime.now().strftime("%d.%m.%Y %H:%M"),
            muted,
        ),
        Paragraph("Ключевые показатели", h2),
        Table(
            [
                ["Показатель", "Значение", "До внедрения AURA"],
                ["Всего обращений", total_pretty, "—"],
                ["Среднее время обработки",
                 f"{summary_data.avg_processing_time} мин",
                 BASELINE["avg_processing_time"]],
                ["Доля решённых", f"{summary_data.resolved_percent}%",
                 BASELINE["resolved_percent"]],
                ["Критических обращений", str(summary_data.critical_count), "—"],
            ],
            colWidths=[70 * mm, 45 * mm, 45 * mm],
            style=table_style,
        ),
        Paragraph("Статистика по регионам", h2),
    ]

    region_data = [["Регион", "Обращений", "Топ-категория", "Тренд"]]
    for r in region_rows:
        sign = "+" if r.trend_percent >= 0 else "−"
        region_data.append([
            r.region, str(r.total), r.top_category,
            f"{sign}{abs(r.trend_percent):.1f}%",
        ])
    story.append(
        Table(region_data, colWidths=[45 * mm, 28 * mm, 62 * mm, 25 * mm],
              style=table_style)
    )

    story.append(Paragraph("Активные всплески", h2))
    if spike_rows:
        for s in spike_rows:
            story.append(Paragraph(
                f"<b>{s.region}</b> · рост {s.percent_increase}% · {s.category}",
                body,
            ))
            story.append(Paragraph(f"{s.description} ({s.time_window})", muted))
            story.append(Spacer(1, 5))
    else:
        story.append(Paragraph("Аномалий не обнаружено.", muted))

    doc.build(story)
    buffer.seek(0)
    name = "aura-pulse-109-" + datetime.now().strftime("%Y%m%d-%H%M") + ".pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="' + name + '"'},
    )


@router.get("/excel")
def export_excel():
    """Excel с листами «Сводка», «По регионам», «Обращения»."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    summary_data = summary()
    region_rows = regions()
    spike_rows = spikes()
    tickets = all_tickets()

    wb = Workbook()
    header_font = Font(bold=True, color="FF6B7280", size=10)
    header_fill = PatternFill("solid", fgColor="FFF9FAFB")

    def style_header(ws, ncols: int, row: int = 1) -> None:
        for col in range(1, ncols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")

    def set_widths(ws, widths: list[int]) -> None:
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width

    # --- Сводка ---
    ws = wb.active
    ws.title = "Сводка"
    ws.append(["Показатель", "Значение", "До внедрения AURA"])
    ws.append(["Всего обращений", summary_data.total_tickets, "—"])
    ws.append(["Среднее время обработки, мин",
               summary_data.avg_processing_time, BASELINE["avg_processing_time"]])
    ws.append(["Доля решённых, %",
               summary_data.resolved_percent, BASELINE["resolved_percent"]])
    ws.append(["Критических обращений", summary_data.critical_count, "—"])
    style_header(ws, 3)

    ws.append([])
    spikes_header_row = ws.max_row + 1
    ws.append(["Регион", "Категория", "Рост, %", "Окно", "Описание"])
    style_header(ws, 5, row=spikes_header_row)
    for s in spike_rows:
        ws.append([s.region, s.category, s.percent_increase,
                   s.time_window, s.description])
    set_widths(ws, [32, 20, 20, 18, 60])

    # --- По регионам ---
    ws = wb.create_sheet("По регионам")
    ws.append(["Регион", "Обращений", "Топ-категория", "Тренд, %"])
    for r in region_rows:
        ws.append([r.region, r.total, r.top_category, r.trend_percent])
    style_header(ws, 4)
    set_widths(ws, [22, 14, 28, 12])
    ws.freeze_panes = "A2"

    # --- Обращения ---
    ws = wb.create_sheet("Обращения")
    ws.append([
        "ID", "Заявитель", "Дата", "Регион", "Источник", "Статус",
        "Приоритет", "Категория", "Подкатегория", "Служба",
        "Уверенность, %", "Требует проверки", "Текст",
    ])
    for t in sorted(tickets, key=lambda x: x["created_at"], reverse=True):
        created = t["created_at"]
        ws.append([
            t["id"],
            # ФИО в системе нет — только номер заявителя
            t.get("citizen_label") or f"Заявитель #{t['id']}",
            created.strftime("%d.%m.%Y %H:%M") if hasattr(created, "strftime") else str(created),
            t["region"], t["source"], t["status"], t["priority"],
            t.get("category") or "—",
            t.get("subcategory") or "—",
            t.get("responsible_service") or "—",
            t.get("confidence_score") or "",
            "да" if t.get("needs_review") else "нет",
            t["text"],
        ])
    style_header(ws, 13)
    set_widths(ws, [6, 14, 17, 13, 10, 12, 11, 18, 22, 22, 15, 17, 90])
    ws.freeze_panes = "A2"
    # Автофильтр по всем колонкам — оператор фильтрует список прямо в Excel
    ws.auto_filter.ref = f"A1:M{ws.max_row}"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    name = "aura-pulse-109-" + datetime.now().strftime("%Y%m%d-%H%M") + ".xlsx"
    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={"Content-Disposition": 'attachment; filename="' + name + '"'},
    )
